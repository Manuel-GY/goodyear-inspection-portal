from django.core.management.base import BaseCommand
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

class Command(BaseCommand):
    help = 'Genera la plantilla Excel oficial de Goodyear (plantilla_carros_goodyear.xlsx) con asignación automática de códigos'

    def add_arguments(self, parser):
        parser.add_argument('--output', type=str, default='plantilla_carros_goodyear.xlsx', help='Ruta donde guardar el archivo Excel generado')

    def handle(self, *args, **options):
        output_path = options.get('output', 'plantilla_carros_goodyear.xlsx')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Plantilla Carros 5S"

        # Cabeceras simplificadas: el sistema autogenera código y nombre automáticamente
        headers = [
            "tipo_carro",
            "turno_o_numero",
            "area",
            "supervisor",
            "ubicacion_especifica",
            "gaveta_1_herramientas",
            "gaveta_2_herramientas",
            "gaveta_3_herramientas",
            "gaveta_4_herramientas",
            "gaveta_5_herramientas",
            "codigo_carro_opcional"
        ]

        ws.append(headers)

        # Estilo de cabecera Goodyear (Navy #0B1D45 y Gold #FBBD00)
        header_fill = PatternFill(start_color="0B1D45", end_color="0B1D45", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FBBD00")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Filas de ejemplo claras
        sample_rows = [
            [
                "TURNO",
                "A",
                "ASRS",
                "Juanito Arias",
                "Pasillo Principal Bahía 1",
                "Juego Llaves Combinadas 8-24mm, Chicharra 1/2\", Dados de Impacto 17-21mm",
                "Destornillador Paleta 6x100mm, Destornillador Cruz PH2, Alicate Universal 8\"",
                "Martillo de Bola 500g, Martillo de Goma, Cincel Plano, Cepillo de Acero",
                "Cinta Métrica 5m, Flexómetro de Trabajo, Manguera Neumática, Pistola de Aire",
                "Candados LOTO Rojos (2 un), Pinza Bloqueo, Tarjeta 5S, Gafas de Seguridad",
                ""  # Deja vacío para autogenerar CH-ASRS-TA
            ],
            [
                "MECANICO",
                "01",
                "Construcción",
                "Juanito Arias",
                "Bahía Mantenimiento Construcción",
                "Juego Dados 1/2\" Heavy Duty (8-32mm), Chicharra Pesada 1/2\", Palanca de Fuerza",
                "Extractor de Rodamientos 3 Patas, Llaves Corona 10-24mm, Llave Ajustable 12\"",
                "Arco de Sierra Profesional, Cinceles Planos, Limas de Ajuste, Llave Stilson 14\"",
                "Pistola Neumática de Impacto 1/2\", Manómetro Digital, Aceite Lubricante",
                "Torquímetro Calibrado 1/2\" (20-200 Nm), Pie de Metro Digital, Kit LOTO",
                ""  # Deja vacío para autogenerar CH-CST-M01
            ],
            [
                "ELECTRICO",
                "01",
                "Final Finish",
                "Juanito Arias",
                "Bahía Eléctrica Final Finish",
                "Juego Destornilladores 1000V VDE (PH0-PH3, SL3-SL6), Pelacables Automático",
                "Multímetro Digital Fluke Calibrado, Pinza Amperimétrica True RMS",
                "Crimpador Terminales Eléctricos, Alicate de Punta VDE, Cautín 60W",
                "Guantes Dieléctricos Clase 0 (1000V), Detector de Tensión Sin Contacto",
                "Candados LOTO Dieléctricos, Tarjetas de Advertencia Eléctrica, Lentes",
                ""  # Deja vacío para autogenerar CH-FF-E01
            ]
        ]

        for row in sample_rows:
            ws.append(row)

        for row in ws.iter_rows(min_row=2, max_row=len(sample_rows)+1):
            for cell in row:
                cell.font = Font(name="Segoe UI", size=10)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center", wrap_text=True)

        column_widths = [16, 16, 20, 20, 28, 42, 42, 42, 42, 42, 24]
        for idx, width in enumerate(column_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = width

        try:
            wb.save(output_path)
            self.stdout.write(self.style.SUCCESS(f"Plantilla Excel generada exitosamente en: {output_path}"))
        except PermissionError:
            alt_path = f"plantilla_carros_goodyear_nueva.xlsx"
            wb.save(alt_path)
            self.stdout.write(self.style.WARNING(f"El archivo '{output_path}' está actualmente abierto en Excel. Se guardó como: {alt_path}"))
