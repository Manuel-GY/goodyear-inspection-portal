from django.test import TestCase, Client
from django.urls import reverse
from django.core.management import call_command
from django.core.files.uploadedfile import SimpleUploadedFile
from inspections.models import ToolCart, DrawerTool, Inspection5S, InspectionMissingItem
import json
import io
import csv
import openpyxl


class ToolCartModelTests(TestCase):
    """Pruebas unitarias para el modelo ToolCart y sus relaciones."""

    def setUp(self):
        self.cart = ToolCart.objects.create(
            codigo_carro="CH-TEST-01",
            nombre_carro="Carro de Prueba 01",
            categoria="TURNO",
            area="Área ASRS",
            supervisor_responsable="Juanito Arias",
            ubicacion_especifica="Bahía 1"
        )

    def test_cart_creation(self):
        """Verifica la correcta creación de un carro de herramientas."""
        self.assertEqual(self.cart.codigo_carro, "CH-TEST-01")
        self.assertEqual(self.cart.nombre_carro, "Carro de Prueba 01")
        self.assertEqual(self.cart.categoria, "TURNO")
        self.assertEqual(self.cart.total_herramientas, 0)
        self.assertEqual(self.cart.estado_general, "OK")
        self.assertIn("CH-TEST-01", str(self.cart))

    def test_drawer_tools_storage_and_recalculation(self):
        """Verifica el guardado de herramientas por gavetas (1 a 5) y recálculo de conteo."""
        # Agregar herramientas en gavetas 1 a 5
        tools_data = [
            (1, "Llave 10mm"),
            (1, "Llave 12mm"),
            (2, "Destornillador PH2"),
            (3, "Martillo de Goma"),
            (4, "Cinta Métrica"),
            (5, "Candado LOTO"),
        ]
        for drawer_num, tool_name in tools_data:
            DrawerTool.objects.create(
                cart=self.cart,
                numero_gaveta=drawer_num,
                nombre_herramienta=tool_name
            )

        self.assertEqual(self.cart.tools.count(), 6)
        count = self.cart.recalculate_tool_count()
        self.assertEqual(count, 6)
        self.assertEqual(self.cart.total_herramientas, 6)
        self.assertEqual(self.cart.estado_general, "OK")

        # Verificar ordenamiento y asignación
        g1_tools = self.cart.tools.filter(numero_gaveta=1)
        self.assertEqual(g1_tools.count(), 2)
        self.assertIn("G1", str(g1_tools.first()))

    def test_cart_deletion_cascade(self):
        """Verifica que al eliminar un carro se eliminan sus herramientas asociadas."""
        DrawerTool.objects.create(cart=self.cart, numero_gaveta=1, nombre_herramienta="Chicharra 1/2")
        self.assertEqual(DrawerTool.objects.filter(cart=self.cart).count(), 1)
        
        cart_db_id = self.cart.id
        self.cart.delete()
        self.assertEqual(DrawerTool.objects.filter(cart_id=cart_db_id).count(), 0)
        self.assertEqual(ToolCart.objects.filter(codigo_carro="CH-TEST-01").count(), 0)


class Inspection5SModelTests(TestCase):
    """Pruebas unitarias para el modelo Inspection5S y cálculo de faltantes."""

    def setUp(self):
        self.cart = ToolCart.objects.create(
            codigo_carro="CH-TEST-02",
            nombre_carro="Carro de Prueba 02",
            categoria="MECANICO",
            area="Área Construcción",
            total_herramientas=10
        )

    def test_inspection_creation_and_missing_items(self):
        """Verifica el registro de inspecciones 5S y desglose de faltantes."""
        insp = Inspection5S.objects.create(
            folio="INS-TEST-001",
            cart=self.cart,
            codigo_carro=self.cart.codigo_carro,
            nombre_auditor="Auditor de Calidad",
            responsable_carro_auditado="Mecánico de Turno",
            supervisor_responsable="Juanito Arias",
            area="Área Construcción",
            estado_dictamen="HERRAMIENTA FALTANTE",
            total_verificadas=8,
            total_herramientas=10,
            detalles_faltantes="Falta Llave Corona 14mm y Chicharra 1/2",
            comentarios_auditor="Se solicita reposición urgente",
            ldap_auditor_id="ac17157",
            ldap_responsable_id="op00123"
        )

        # Crear faltantes
        InspectionMissingItem.objects.create(
            inspection=insp,
            numero_gaveta=1,
            nombre_herramienta_faltante="Llave Corona 14mm"
        )
        InspectionMissingItem.objects.create(
            inspection=insp,
            numero_gaveta=1,
            nombre_herramienta_faltante="Chicharra 1/2"
        )

        self.assertEqual(insp.folio, "INS-TEST-001")
        self.assertEqual(insp.missing_items.count(), 2)
        self.assertEqual(insp.total_herramientas - insp.total_verificadas, 2)
        self.assertIn("INS-TEST-001", str(insp))


class ApiRestEndpointsTests(TestCase):
    """Pruebas exhaustivas para todos los endpoints de la API REST."""

    def setUp(self):
        self.client = Client()
        self.cart = ToolCart.objects.create(
            codigo_carro="CH-API-01",
            nombre_carro="Carro API Test",
            categoria="ELECTRICO",
            area="Área Final Finish",
            supervisor_responsable="Juanito Arias",
            ubicacion_especifica="Bahía Eléctrica"
        )
        # 3 herramientas
        DrawerTool.objects.create(cart=self.cart, numero_gaveta=1, nombre_herramienta="Multímetro", orden_posicion=1)
        DrawerTool.objects.create(cart=self.cart, numero_gaveta=2, nombre_herramienta="Alicate VDE", orden_posicion=1)
        DrawerTool.objects.create(cart=self.cart, numero_gaveta=5, nombre_herramienta="Guantes Dieléctricos", orden_posicion=1)
        self.cart.recalculate_tool_count()

    def test_index_view(self):
        """Verifica que la vista principal renderice el template index.html con status 200."""
        response = self.client.get(reverse('index'))
        self.assertEqual(response.status_code, 200)

    def test_logo_view(self):
        """Verifica que el endpoint de logo retorne la imagen PNG."""
        response = self.client.get(reverse('logo_view'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'image/png')

    def test_api_drawer_structure(self):
        """Verifica que el endpoint de estructura devuelva las categorías 5S estándar."""
        response = self.client.get(reverse('api_drawer_structure'))
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['status'], 'success')
        self.assertIn('TURNO', data['drawers'])
        self.assertIn('MECANICO', data['drawers'])
        self.assertIn('ELECTRICO', data['drawers'])
        self.assertIn('MECATRONICO', data['drawers'])

    def test_api_carts_get_and_post(self):
        """Verifica GET y POST en /api/carts/."""
        # GET
        response = self.client.get(reverse('api_carts_list_create'))
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['status'], 'success')
        self.assertIn('CH-API-01', data['carts'])
        self.assertEqual(data['carts']['CH-API-01']['toolsCount'], 3)

        # POST (creación válida)
        new_cart_payload = {
            "codigo_carro": "CH-NEW-01",
            "nombre_carro": "Carro Nuevo Automatización",
            "categoria": "MECATRONICO",
            "area": "Área ASRS",
            "supervisor_responsable": "Juanito Arias",
            "ubicacion_especifica": "Línea 2"
        }
        res_post = self.client.post(
            reverse('api_carts_list_create'),
            data=json.dumps(new_cart_payload),
            content_type='application/json'
        )
        self.assertEqual(res_post.status_code, 201)
        res_data = res_post.json()
        self.assertEqual(res_data['status'], 'success')
        self.assertTrue(ToolCart.objects.filter(codigo_carro="CH-NEW-01").exists())

        # POST (duplicado -> error 400)
        res_dup = self.client.post(
            reverse('api_carts_list_create'),
            data=json.dumps(new_cart_payload),
            content_type='application/json'
        )
        self.assertEqual(res_dup.status_code, 400)

        # POST (datos incompletos -> error 400)
        res_empty = self.client.post(
            reverse('api_carts_list_create'),
            data=json.dumps({"codigo_carro": ""}),
            content_type='application/json'
        )
        self.assertEqual(res_empty.status_code, 400)

    def test_api_cart_detail_update_delete(self):
        """Verifica GET, PUT y DELETE en /api/carts/<cart_id>/."""
        # GET detalle
        url = reverse('api_cart_detail_update_delete', kwargs={'cart_id': 'CH-API-01'})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['status'], 'success')
        self.assertEqual(data['cart']['codigo_carro'], 'CH-API-01')
        self.assertEqual(len(data['cart']['drawers']['1']), 1)
        self.assertIn("Multímetro", data['cart']['drawers']['1'])

        # PUT actualización
        update_payload = {
            "nombre_carro": "Carro Eléctrico Modificado",
            "ubicacion_especifica": "Bahía Central 5S",
            "estado_general": "Observado"
        }
        res_put = self.client.put(url, data=json.dumps(update_payload), content_type='application/json')
        self.assertEqual(res_put.status_code, 200)
        self.cart.refresh_from_db()
        self.assertEqual(self.cart.nombre_carro, "Carro Eléctrico Modificado")
        self.assertEqual(self.cart.ubicacion_especifica, "Bahía Central 5S")
        self.assertEqual(self.cart.estado_general, "Observado")

        # DELETE
        res_del = self.client.delete(url)
        self.assertEqual(res_del.status_code, 200)
        self.assertFalse(ToolCart.objects.filter(codigo_carro="CH-API-01").exists())

        # GET 404
        res_not_found = self.client.get(url)
        self.assertEqual(res_not_found.status_code, 404)

    def test_api_cart_drawer_tools(self):
        """Verifica POST en /api/carts/<cart_id>/drawers/<drawer_num>/tools/."""
        url = reverse('api_cart_drawer_tools', kwargs={'cart_id': 'CH-API-01', 'drawer_num': 3})
        payload = {
            "tools": [
                "Martillo de Bola 500g",
                "Cincel Plano de Ajuste",
                "Cepillo de Alambre"
            ]
        }
        response = self.client.post(url, data=json.dumps(payload), content_type='application/json')
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['status'], 'success')
        
        # Verificar en base de datos
        self.cart.refresh_from_db()
        g3_tools = self.cart.tools.filter(numero_gaveta=3).order_by('orden_posicion')
        self.assertEqual(g3_tools.count(), 3)
        self.assertEqual(g3_tools.first().nombre_herramienta, "Martillo de Bola 500g")
        self.assertEqual(self.cart.total_herramientas, 6) # 3 iniciales + 3 nuevas

        # Test gaveta inválida (e.g. 6 -> error 400)
        bad_url = reverse('api_cart_drawer_tools', kwargs={'cart_id': 'CH-API-01', 'drawer_num': 6})
        bad_res = self.client.post(bad_url, data=json.dumps(payload), content_type='application/json')
        self.assertEqual(bad_res.status_code, 400)

    def test_api_inspections_list_and_create(self):
        """Verifica listado y registro de inspecciones en /api/inspections/."""
        payload = {
            "folio": "5S-2026-0001",
            "cartId": "CH-API-01",
            "cartArea": "Área Final Finish",
            "auditor": "Juan Pérez",
            "responsible": "Carlos González",
            "supervisor": "Juanito Arias",
            "status": "HERRAMIENTA FALTANTE",
            "totalChecked": 2,
            "totalTools": 3,
            "missingDetails": "Falta Guantes Dieléctricos",
            "comments": "Inspección de turno mañana",
            "ldapAuditor": "jperez12",
            "ldapResp": "cgonzalez34",
            "missingItemsList": [
                {"drawer": 5, "tool": "Guantes Dieléctricos"}
            ]
        }
        res_post = self.client.post(
            reverse('api_inspections_list_create'),
            data=json.dumps(payload),
            content_type='application/json'
        )
        self.assertEqual(res_post.status_code, 201)
        data = res_post.json()
        self.assertEqual(data['status'], 'success')
        self.assertEqual(data['folio'], '5S-2026-0001')

        # GET inspecciones
        res_get = self.client.get(reverse('api_inspections_list_create'))
        self.assertEqual(res_get.status_code, 200)
        get_data = res_get.json()
        self.assertEqual(get_data['status'], 'success')
        self.assertEqual(get_data['count'], 1)
        first_insp = get_data['inspections'][0]
        self.assertEqual(first_insp['folio'], '5S-2026-0001')
        self.assertEqual(first_insp['missingCount'], 1)

    def test_api_dashboard_stats(self):
        """Verifica el cálculo de KPIs y estadísticas en /api/dashboard/stats/."""
        # Registrar una inspección conforme y una no conforme
        Inspection5S.objects.create(
            folio="INS-CONF-01",
            cart=self.cart,
            codigo_carro="CH-API-01",
            nombre_auditor="Inspector",
            responsable_carro_auditado="Mecánico",
            area="Área Final Finish",
            estado_dictamen="CONFORME 100%",
            total_verificadas=3,
            total_herramientas=3
        )
        Inspection5S.objects.create(
            folio="INS-MISS-01",
            cart=self.cart,
            codigo_carro="CH-API-01",
            nombre_auditor="Inspector",
            responsable_carro_auditado="Mecánico",
            area="Área Final Finish",
            estado_dictamen="HERRAMIENTA FALTANTE",
            total_verificadas=2,
            total_herramientas=3
        )

        response = self.client.get(reverse('api_dashboard_stats'))
        self.assertEqual(response.status_code, 200)
        stats = response.json()
        self.assertEqual(stats['status'], 'success')
        self.assertEqual(stats['totalCarts'], 1)
        self.assertEqual(stats['totalAudits'], 2)
        self.assertEqual(stats['compliantAudits'], 1)
        self.assertEqual(stats['missingAudits'], 1)
        self.assertEqual(stats['passRate'], 50)
        self.assertIn('Área Final Finish', stats['areas'])


class ExcelAndCSVImportExportTests(TestCase):
    """Pruebas para generación de plantilla y procesamiento de importaciones."""

    def setUp(self):
        self.client = Client()

    def test_download_excel_template(self):
        """Verifica la generación y descarga de la plantilla Excel oficial."""
        response = self.client.get(reverse('api_download_excel_template'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('plantilla_carros_goodyear.xlsx', response['Content-Disposition'])

        # Validar contenido openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        self.assertEqual(ws.title, "Plantilla Carros 5S")
        headers = [cell.value for cell in ws[1]]
        self.assertIn("codigo_carro", headers)
        self.assertIn("nombre_carro", headers)
        self.assertIn("gaveta_1_herramientas", headers)

    def test_import_csv_api(self):
        """Verifica la importación masiva de carros y gavetas desde archivo CSV."""
        csv_content = (
            "codigo_carro,nombre_carro,categoria,area,supervisor,ubicacion_especifica,gaveta_1_herramientas,gaveta_2_herramientas,gaveta_3_herramientas,gaveta_4_herramientas,gaveta_5_herramientas\n"
            "CH-CSV-01,Carro CSV Test,MECANICO,Área ASRS,Juanito Arias,Pasillo 4,\"Llave 10, Llave 12\",\"Destornillador PH1, Destornillador PH2\",\"Martillo\",\"Flexometro\",\"Candado LOTO\"\n"
        )
        csv_file = SimpleUploadedFile("carros_test.csv", csv_content.encode('utf-8-sig'), content_type="text/csv")

        response = self.client.post(reverse('api_import_excel'), {'file': csv_file}, format='multipart')
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['status'], 'success')
        self.assertEqual(data['imported_carts'], 1)
        self.assertEqual(data['imported_tools'], 7)

        # Verificar en base de datos
        cart = ToolCart.objects.filter(codigo_carro="CH-CSV-01").first()
        self.assertIsNotNone(cart)
        self.assertEqual(cart.nombre_carro, "Carro CSV Test")
        self.assertEqual(cart.total_herramientas, 7)
        self.assertEqual(cart.tools.filter(numero_gaveta=1).count(), 2)
        self.assertEqual(cart.tools.filter(numero_gaveta=2).count(), 2)

    def test_import_excel_api(self):
        """Verifica la importación masiva de carros desde archivo Excel (.xlsx)."""
        wb = openpyxl.Workbook()
        ws = wb.active
        headers = ["codigo_carro", "nombre_carro", "categoria", "area", "supervisor", "ubicacion_especifica", "gaveta_1_herramientas"]
        ws.append(headers)
        ws.append(["CH-XLSX-01", "Carro XLSX Test", "TURNO", "Área Construcción", "Juanito Arias", "Bahía 2", "Llave Francesa; Alicate de Punta"])
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        excel_file = SimpleUploadedFile("carros_test.xlsx", buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response = self.client.post(reverse('api_import_excel'), {'file': excel_file}, format='multipart')
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['status'], 'success')
        self.assertEqual(data['imported_carts'], 1)
        self.assertEqual(data['imported_tools'], 2)

        cart = ToolCart.objects.filter(codigo_carro="CH-XLSX-01").first()
        self.assertIsNotNone(cart)
        self.assertEqual(cart.total_herramientas, 2)


class ManagementCommandsTests(TestCase):
    """Pruebas para los comandos de gestión Django."""

    def test_seed_data_and_reset_factory(self):
        """Verifica la ejecución del comando seed_data y api_reset_factory."""
        ToolCart.objects.all().delete()
        self.assertEqual(ToolCart.objects.count(), 0)

        # Ejecutar seed_data
        call_command('seed_data')
        self.assertEqual(ToolCart.objects.count(), 28)
        
        # Carro con herramientas oficiales cargadas
        asrs_ta = ToolCart.objects.get(codigo_carro="CH-ASRS-TA")
        self.assertGreater(asrs_ta.total_herramientas, 0)
        self.assertEqual(asrs_ta.estado_general, "OK")

        # Probar endpoint api_reset_factory
        client = Client()
        res = client.post(reverse('api_reset_factory'))
        self.assertEqual(res.status_code, 200)
        self.assertEqual(ToolCart.objects.count(), 28)

    def test_export_template_command(self):
        """Verifica la ejecución del comando export_template."""
        call_command('export_template')
        import os
        self.assertTrue(os.path.exists("plantilla_carros_goodyear.xlsx"))
